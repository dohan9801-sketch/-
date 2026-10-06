"""input 폴더의 지점별 판매 엑셀을 모두 합쳐 월간 보고서를 만듭니다.

실행하면 output 폴더에 다음 시트가 들어 있는 보고서 엑셀이 생깁니다.
  요약       : 총매출·지점 순위·분류별 매출 + 차트
  지점x분류  : 지점별·분류별 매출 표
  일별추이   : 날짜별 지점 매출 + 추이 차트
  상품순위   : 상품별 매출 순위
  전체데이터 : 모든 지점의 판매 내역을 한 표로 합친 원본

지점 이름은 파일 이름의 "_" 앞부분에서 가져옵니다. (예: 강남점_2026년09월.xlsx → 강남점)
파일마다 제목 줄이 있거나, 열 순서가 다르거나, 중간에 빈 줄이 있어도 처리합니다.

사용법:
  python merge_report.py
  python merge_report.py --input 다른폴더 --output 저장폴더
"""
import argparse
import sys
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent

REQUIRED = ["날짜", "상품분류", "상품명", "수량", "단가", "금액"]

DARK = PatternFill("solid", fgColor="1F2937")
LIGHT = PatternFill("solid", fgColor="F3F4F6")
ACCENT = PatternFill("solid", fgColor="E8F0FE")
THIN = Side(style="thin", color="D1D5DB")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
MONEY = "#,##0"
PCT = "0.0%"


# ---------------------------------------------------------------- 읽기

def find_header(ws):
    """'날짜'와 '금액'이 함께 있는 줄을 머리글로 봅니다. (앞쪽 10줄 안에서 찾음)"""
    for row_no, row in enumerate(ws.iter_rows(max_row=10, values_only=True), start=1):
        names = [str(v).strip() if v is not None else "" for v in row]
        if "날짜" in names and "금액" in names:
            return row_no, names
    return None, None


def to_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip().replace(".", "-").replace("/", "-")
    return datetime.strptime(text[:10], "%Y-%m-%d").date()


def to_number(value):
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, (int, float)):
        return value
    return float(str(value).replace(",", "").replace("원", "").strip())


def read_branch_file(path):
    """파일 하나를 읽어 판매 내역 목록과 경고 목록을 돌려줍니다."""
    branch = path.stem.split("_")[0]
    ws = load_workbook(path, data_only=True).active
    header_row, names = find_header(ws)
    if header_row is None:
        return branch, [], [f"{path.name}: '날짜', '금액' 머리글을 찾지 못해 건너뜀"]
    missing = [c for c in REQUIRED if c not in names]
    if missing:
        return branch, [], [f"{path.name}: {', '.join(missing)} 열이 없어 건너뜀"]
    col = {c: names.index(c) for c in REQUIRED}

    records, warnings = [], []
    for row_no, row in enumerate(ws.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
        if all(v is None or str(v).strip() == "" for v in row):
            continue  # 빈 줄
        get = lambda c: row[col[c]] if col[c] < len(row) else None  # noqa: E731
        try:
            qty = to_number(get("수량")) or 0
            price = to_number(get("단가")) or 0
            amount = to_number(get("금액"))
            records.append({
                "지점": branch,
                "날짜": to_date(get("날짜")),
                "상품분류": str(get("상품분류") or "미분류").strip(),
                "상품명": str(get("상품명") or "").strip(),
                "수량": qty,
                "단가": price,
                "금액": amount if amount is not None else qty * price,
            })
        except (ValueError, TypeError):
            warnings.append(f"{path.name} {row_no}행: 값을 읽을 수 없어 건너뜀 {list(row)}")
    return branch, records, warnings


# ---------------------------------------------------------------- 서식 도우미

def write_table(ws, top, left, headers, rows, formats=None, total=None):
    """머리글과 테두리가 있는 표를 쓰고, 마지막 줄 번호를 돌려줍니다."""
    formats = formats or {}
    for j, h in enumerate(headers):
        c = ws.cell(row=top, column=left + j, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = DARK
        c.alignment = Alignment(horizontal="center")
        c.border = BOX
    r = top
    for r_offset, values in enumerate(rows, start=1):
        r = top + r_offset
        for j, v in enumerate(values):
            c = ws.cell(row=r, column=left + j, value=v)
            c.border = BOX
            if j in formats:
                c.number_format = formats[j]
    if total:
        r += 1
        for j, v in enumerate(total):
            c = ws.cell(row=r, column=left + j, value=v)
            c.font = Font(bold=True)
            c.fill = LIGHT
            c.border = BOX
            if j in formats:
                c.number_format = formats[j]
    return r


def set_widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def title(ws, text, sub):
    ws["A1"] = text
    ws["A1"].font = Font(bold=True, size=16)
    ws["A2"] = sub
    ws["A2"].font = Font(color="6B7280")


# ---------------------------------------------------------------- 보고서

def build_report(records, period, files, out_path):
    branches = sorted({r["지점"] for r in records})
    categories = sorted({r["상품분류"] for r in records})
    days = sorted({r["날짜"] for r in records})

    by_branch = defaultdict(lambda: [0, 0])  # [매출, 수량]
    by_cat = defaultdict(int)
    by_branch_cat = defaultdict(int)
    by_day_branch = defaultdict(int)
    by_product = defaultdict(lambda: [None, 0, 0])  # [분류, 수량, 매출]
    for r in records:
        by_branch[r["지점"]][0] += r["금액"]
        by_branch[r["지점"]][1] += r["수량"]
        by_cat[r["상품분류"]] += r["금액"]
        by_branch_cat[(r["지점"], r["상품분류"])] += r["금액"]
        by_day_branch[(r["날짜"], r["지점"])] += r["금액"]
        p = by_product[r["상품명"]]
        p[0] = r["상품분류"]
        p[1] += r["수량"]
        p[2] += r["금액"]

    total_sales = sum(v[0] for v in by_branch.values())
    total_qty = sum(v[1] for v in by_branch.values())
    ranking = sorted(by_branch.items(), key=lambda kv: kv[1][0], reverse=True)
    created = datetime.now().strftime("%Y-%m-%d %H:%M")

    wb = Workbook()

    # 1) 요약
    ws = wb.active
    ws.title = "요약"
    title(ws, f"{period} 지점별 판매 보고서", f"생성: {created} · 취합 파일 {len(files)}개 · 판매 내역 {len(records):,}건")

    kpis = [("총매출", total_sales, MONEY + '"원"'), ("총 판매수량", total_qty, MONEY + '"개"'),
            ("지점 수", len(branches), '0"곳"'), ("매출 1위", ranking[0][0], "@")]
    for i, (label, value, fmt) in enumerate(kpis):
        col = 1 + i * 2
        ws.merge_cells(start_row=4, start_column=col, end_row=4, end_column=col + 1)
        ws.merge_cells(start_row=5, start_column=col, end_row=5, end_column=col + 1)
        lc = ws.cell(row=4, column=col, value=label)
        lc.font = Font(color="6B7280")
        vc = ws.cell(row=5, column=col, value=value)
        vc.font = Font(bold=True, size=15, color="1F56C4")
        vc.number_format = fmt
        for r in (4, 5):
            for c in (col, col + 1):
                ws.cell(row=r, column=c).fill = ACCENT
                ws.cell(row=r, column=c).alignment = Alignment(horizontal="center")

    ws["A7"] = "지점별 매출 순위"
    ws["A7"].font = Font(bold=True, size=12)
    rows = [[i, b, v[0], v[1], v[0] / total_sales] for i, (b, v) in enumerate(ranking, start=1)]
    end = write_table(ws, 8, 1, ["순위", "지점", "매출", "수량", "비중"], rows,
                      {2: MONEY, 3: MONEY, 4: PCT}, ["합계", "", total_sales, total_qty, 1])
    branch_first, branch_last = 9, 8 + len(rows)

    cat_top = end + 3
    ws.cell(row=cat_top - 1, column=1, value="분류별 매출").font = Font(bold=True, size=12)
    cat_rows = sorted(([c, by_cat[c], by_cat[c] / total_sales] for c in categories), key=lambda x: -x[1])
    write_table(ws, cat_top, 1, ["분류", "매출", "비중"], cat_rows, {1: MONEY, 2: PCT})

    bar = BarChart()
    bar.title = "지점별 매출"
    bar.y_axis.title = "매출(원)"
    bar.y_axis.numFmt = MONEY
    bar.legend = None
    bar.add_data(Reference(ws, min_col=3, min_row=8, max_row=branch_last), titles_from_data=True)
    bar.set_categories(Reference(ws, min_col=2, min_row=branch_first, max_row=branch_last))
    bar.height, bar.width = 7.5, 14
    ws.add_chart(bar, "H7")

    pie = PieChart()
    pie.title = "분류별 매출 비중"
    pie.add_data(Reference(ws, min_col=2, min_row=cat_top, max_row=cat_top + len(cat_rows)), titles_from_data=True)
    pie.set_categories(Reference(ws, min_col=1, min_row=cat_top + 1, max_row=cat_top + len(cat_rows)))
    pie.height, pie.width = 7.5, 14
    ws.add_chart(pie, "H23")
    set_widths(ws, [8, 12, 14, 10, 9, 4, 4])

    # 2) 지점 x 분류
    ws = wb.create_sheet("지점x분류")
    title(ws, "지점별 · 분류별 매출", period)
    rows = [[b] + [by_branch_cat[(b, c)] for c in categories] + [by_branch[b][0]] for b in branches]
    totals = ["합계"] + [by_cat[c] for c in categories] + [total_sales]
    fmts = {j: MONEY for j in range(1, len(categories) + 2)}
    write_table(ws, 4, 1, ["지점"] + categories + ["합계"], rows, fmts, totals)
    set_widths(ws, [12] + [14] * (len(categories) + 1))

    # 3) 일별 추이
    ws = wb.create_sheet("일별추이")
    title(ws, "일별 매출 추이", period)
    rows = []
    for d in days:
        values = [by_day_branch[(d, b)] for b in branches]
        rows.append([d] + values + [sum(values)])
    fmts = {0: "yyyy-mm-dd"} | {j: MONEY for j in range(1, len(branches) + 2)}
    totals = ["합계"] + [by_branch[b][0] for b in branches] + [total_sales]
    write_table(ws, 4, 1, ["날짜"] + branches + ["합계"], rows, fmts, totals)
    ws.freeze_panes = "B5"
    set_widths(ws, [13] + [12] * (len(branches) + 1))

    line = LineChart()
    line.title = "일별 총매출"
    line.y_axis.numFmt = MONEY
    line.legend = None
    total_col = len(branches) + 2
    line.add_data(Reference(ws, min_col=total_col, min_row=4, max_row=4 + len(days)), titles_from_data=True)
    line.set_categories(Reference(ws, min_col=1, min_row=5, max_row=4 + len(days)))
    line.height, line.width = 8, 18
    ws.add_chart(line, f"{get_column_letter(total_col + 2)}4")

    # 4) 상품 순위
    ws = wb.create_sheet("상품순위")
    title(ws, "상품별 매출 순위", period)
    products = sorted(by_product.items(), key=lambda kv: kv[1][2], reverse=True)
    rows = [[i, name, v[0], v[1], v[2], v[2] / total_sales] for i, (name, v) in enumerate(products, start=1)]
    write_table(ws, 4, 1, ["순위", "상품명", "분류", "수량", "매출", "비중"], rows, {3: MONEY, 4: MONEY, 5: PCT})
    set_widths(ws, [7, 18, 10, 10, 14, 9])

    # 5) 전체 데이터
    ws = wb.create_sheet("전체데이터")
    headers = ["지점"] + REQUIRED
    rows = [[r[h] for h in headers] for r in sorted(records, key=lambda r: (r["날짜"], r["지점"]))]
    write_table(ws, 1, 1, headers, rows, {1: "yyyy-mm-dd", 4: MONEY, 5: MONEY, 6: MONEY})
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
    ws.freeze_panes = "A2"
    set_widths(ws, [10, 12, 10, 16, 8, 10, 12])

    # 인쇄할 때 가로 방향, 한 페이지 너비에 맞춤
    for sheet in wb.worksheets:
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)


def main():
    parser = argparse.ArgumentParser(description="지점별 판매 엑셀을 합쳐 월간 보고서 생성")
    parser.add_argument("--input", type=Path, default=HERE / "input", help="지점 파일이 있는 폴더")
    parser.add_argument("--output", type=Path, default=HERE / "output", help="보고서를 저장할 폴더")
    args = parser.parse_args()

    # ~$ 로 시작하는 파일은 엑셀이 열려 있을 때 생기는 임시 파일이라 제외
    files = sorted(p for p in args.input.glob("*.xlsx") if not p.name.startswith("~$"))
    if not files:
        sys.exit(f"{args.input} 폴더에 엑셀 파일이 없습니다.")

    records, warnings = [], []
    for path in files:
        branch, rows, warn = read_branch_file(path)
        records += rows
        warnings += warn
        print(f"읽음: {path.name:<24} {branch} {len(rows):>5,}건")

    if not records:
        sys.exit("읽은 판매 내역이 없습니다.")

    months = sorted({(r["날짜"].year, r["날짜"].month) for r in records})
    if len(months) == 1:
        period = f"{months[0][0]}년 {months[0][1]:02d}월"
        stamp = f"{months[0][0]}-{months[0][1]:02d}"
    else:
        period = f"{months[0][0]}년 {months[0][1]:02d}월 ~ {months[-1][0]}년 {months[-1][1]:02d}월"
        stamp = f"{months[0][0]}-{months[0][1]:02d}_{months[-1][0]}-{months[-1][1]:02d}"

    out = args.output / f"월간보고서_{stamp}.xlsx"
    try:
        build_report(records, period, files, out)
    except PermissionError:
        sys.exit(f"저장할 수 없습니다. {out.name} 파일이 엑셀에서 열려 있으면 닫고 다시 실행하세요.")

    for w in warnings:
        print("주의:", w)
    print(f"\n총 {len(files)}개 파일, {len(records):,}건을 합쳤습니다.")
    print(f"보고서: {out}")


if __name__ == "__main__":
    main()
