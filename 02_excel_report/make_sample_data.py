"""시연용 지점별 판매 엑셀(input/*.xlsx)을 만듭니다.

실제 현장처럼 지점마다 파일 모양이 조금씩 다르게 만들었습니다.
  - 부산점: 맨 위에 제목 줄이 두 줄 있음
  - 대구점: 열 순서가 다름
  - 광주점: 중간에 빈 줄이 있음
"""
import random
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

HERE = Path(__file__).resolve().parent
INPUT_DIR = HERE / "input"

YEAR, MONTH = 2026, 9

BRANCHES = {
    "강남점": 1.4,
    "홍대점": 1.2,
    "부산점": 1.0,
    "대구점": 0.8,
    "광주점": 0.7,
    "대전점": 0.6,
}

PRODUCTS = {
    "음료": [("아메리카노", 4500), ("카페라떼", 5000), ("바닐라라떼", 5500), ("자몽에이드", 6000)],
    "디저트": [("치즈케이크", 6500), ("크루아상", 4000), ("마카롱", 3000)],
    "상품": [("원두 200g", 15000), ("텀블러", 22000), ("드립백 10개입", 12000)],
}

HEADERS = ["날짜", "상품분류", "상품명", "수량", "단가", "금액"]


def sales_rows(weight, rng):
    rows = []
    day = date(YEAR, MONTH, 1)
    while day.month == MONTH:
        weekend = day.weekday() >= 5
        for category, items in PRODUCTS.items():
            for name, price in items:
                if rng.random() < 0.35:
                    continue  # 매일 모든 상품이 팔리지는 않음
                base = 12 if category == "음료" else 5 if category == "디저트" else 1
                qty = max(1, round(rng.uniform(0.5, 1.5) * base * weight * (1.3 if weekend else 1)))
                rows.append([day, category, name, qty, price, qty * price])
        day += timedelta(days=1)
    return rows


def style_header(ws, row):
    for cell in ws[row]:
        if cell.value is not None:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="E5E7EB")


def main():
    rng = random.Random(42)  # 매번 같은 데이터가 나오도록 고정
    INPUT_DIR.mkdir(exist_ok=True)

    for branch, weight in BRANCHES.items():
        rows = sales_rows(weight, rng)
        wb = Workbook()
        ws = wb.active
        ws.title = "판매내역"

        if branch == "부산점":
            ws.append([f"{branch} {MONTH}월 판매 내역"])
            ws["A1"].font = Font(bold=True, size=14)
            ws.append(["작성자: 점장"])
            header_row = 3
            ws.append(HEADERS)
            for r in rows:
                ws.append(r)
        elif branch == "대구점":
            order = [2, 0, 1, 3, 4, 5]  # 상품명, 날짜, 상품분류, 수량, 단가, 금액
            header_row = 1
            ws.append([HEADERS[i] for i in order])
            for r in rows:
                ws.append([r[i] for i in order])
        else:
            header_row = 1
            ws.append(HEADERS)
            for i, r in enumerate(rows):
                if branch == "광주점" and i in (40, 41, 120):
                    ws.append([])
                ws.append(r)

        style_header(ws, header_row)
        for row in ws.iter_rows(min_row=header_row + 1):
            for cell in row:
                if hasattr(cell.value, "year"):
                    cell.number_format = "yyyy-mm-dd"
                elif isinstance(cell.value, int):
                    cell.number_format = "#,##0"
        for col, width in zip("ABCDEF", [12, 10, 14, 8, 10, 12]):
            ws.column_dimensions[col].width = width

        path = INPUT_DIR / f"{branch}_{YEAR}년{MONTH:02d}월.xlsx"
        wb.save(path)
        print(f"만들었습니다: {path.name} ({len(rows)}건)")


if __name__ == "__main__":
    main()
