"""엑셀(data/상품목록.xlsx)의 상품을 관리자 페이지에 자동으로 등록합니다.

실행하면
  1) 브라우저를 열고 관리자 페이지에 로그인한 뒤
  2) 엑셀의 상품을 한 줄씩 입력·등록하고
  3) 줄마다 처리 결과(완료/실패 사유)를 적은 결과 엑셀을 output 폴더에 저장합니다.

사용법:
  python auto_input.py              # 화면을 보면서 천천히 실행 (시연용)
  python auto_input.py --fast       # 빠르게 실행
  python auto_input.py --headless   # 브라우저 창 없이 실행
  python auto_input.py --excel 다른파일.xlsx
"""
import argparse
import os
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from playwright.sync_api import Error as PlaywrightError
from playwright.sync_api import sync_playwright

HERE = Path(__file__).resolve().parent
DEFAULT_EXCEL = HERE / "data" / "상품목록.xlsx"
DEFAULT_URL = (HERE / "demo_site" / "index.html").as_uri()
OUTPUT_DIR = HERE / "output"

LOGIN_ID = "admin"
LOGIN_PW = "1234"

# 엑셀 열 이름 -> 프로그램에서 쓰는 이름
COLUMNS = {
    "상품명": "name",
    "카테고리": "category",
    "판매가": "price",
    "재고": "stock",
    "배송방식": "shipping",
    "즉시판매": "on_sale",
    "상품설명": "desc",
}

OK_FILL = PatternFill("solid", fgColor="DCFCE7")
FAIL_FILL = PatternFill("solid", fgColor="FEE2E2")


def read_products(ws):
    """엑셀 시트에서 (행 번호, 상품 정보) 목록을 읽습니다."""
    headers = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
    missing = [name for name in COLUMNS if name not in headers]
    if missing:
        sys.exit(f"엑셀 첫 줄에 다음 열이 없습니다: {', '.join(missing)}")
    index = {COLUMNS[name]: headers.index(name) for name in COLUMNS}

    products = []
    for row_no, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if all(v is None or str(v).strip() == "" for v in row):
            continue  # 빈 줄 건너뛰기
        item = {}
        for key, col in index.items():
            value = row[col] if col < len(row) else None
            item[key] = "" if value is None else str(value).strip()
        products.append((row_no, item))
    return products


def text_number(value):
    """엑셀에서 19900.0 처럼 읽힌 숫자를 19900 으로 정리합니다."""
    try:
        number = float(value.replace(",", ""))
    except ValueError:
        return value
    return str(int(number)) if number.is_integer() else str(number)


def launch_browser(p, headless, slow_mo):
    """설치된 크롬 → 엣지 → Playwright 내장 크로미움 순서로 브라우저를 엽니다."""
    options = {"headless": headless, "slow_mo": slow_mo}
    custom = os.environ.get("BROWSER_PATH")
    if custom:
        return p.chromium.launch(executable_path=custom, **options)
    for channel in ("chrome", "msedge"):
        try:
            return p.chromium.launch(channel=channel, **options)
        except PlaywrightError:
            continue
    return p.chromium.launch(**options)


def login(page, url):
    page.goto(url)
    page.fill("#login-id", LOGIN_ID)
    page.fill("#login-pw", LOGIN_PW)
    page.click("#login-btn")
    page.wait_for_selector("#product-form", state="visible", timeout=5000)


def register(page, item):
    """상품 하나를 입력하고 등록합니다. 실패하면 사유를 담은 예외를 냅니다."""
    before = int(page.text_content("#count") or 0)

    page.fill("#name", item["name"])
    if item["category"]:
        page.select_option("#category", label=item["category"])
    else:
        page.select_option("#category", value="")
    page.fill("#price", text_number(item["price"]))
    page.fill("#stock", text_number(item["stock"]))
    if item["shipping"]:
        page.check(f"input[name=shipping][value='{item['shipping']}']")
    page.set_checked("#on-sale", item["on_sale"].upper() in ("Y", "O", "예", "TRUE", "1"))
    page.fill("#desc", item["desc"])
    page.click("#submit-btn")

    # 등록되면 목록 개수가 늘어나고, 실패하면 화면에 오류 문구가 뜹니다.
    page.wait_for_function(
        "([before]) => Number(document.querySelector('#count').textContent) > before"
        " || document.querySelector('#form-error').textContent.trim() !== ''",
        arg=[before],
        timeout=5000,
    )
    error = (page.text_content("#form-error") or "").strip()
    if error:
        raise RuntimeError(error)


def main():
    parser = argparse.ArgumentParser(description="엑셀 상품 목록을 관리자 페이지에 자동 등록")
    parser.add_argument("--excel", type=Path, default=DEFAULT_EXCEL, help="입력 엑셀 파일")
    parser.add_argument("--url", default=DEFAULT_URL, help="관리자 페이지 주소")
    parser.add_argument("--fast", action="store_true", help="입력 속도를 늦추지 않음")
    parser.add_argument("--headless", action="store_true", help="브라우저 창을 띄우지 않음")
    args = parser.parse_args()

    if not args.excel.exists():
        sys.exit(f"엑셀 파일을 찾을 수 없습니다: {args.excel}")

    wb = load_workbook(args.excel)
    ws = wb.active
    products = read_products(ws)
    if not products:
        sys.exit("엑셀에 등록할 상품이 없습니다.")

    # 결과를 적을 열 추가
    result_col = ws.max_column + 1
    time_col = result_col + 1
    for col, title in ((result_col, "처리결과"), (time_col, "처리시각")):
        cell = ws.cell(row=1, column=col, value=title)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F2937")
    ws.column_dimensions[ws.cell(row=1, column=result_col).column_letter].width = 30
    ws.column_dimensions[ws.cell(row=1, column=time_col).column_letter].width = 20

    print(f"총 {len(products)}건을 등록합니다.\n")
    ok = fail = 0
    # 시연할 때 입력 과정이 눈에 보이도록 동작마다 잠깐 쉽니다.
    slow_mo = 0 if args.fast else 120

    with sync_playwright() as p:
        browser = launch_browser(p, args.headless, slow_mo)
        page = browser.new_page(viewport={"width": 1280, "height": 800})
        page.set_default_timeout(5000)
        login(page, args.url)

        for i, (row_no, item) in enumerate(products, start=1):
            label = item["name"] or f"{row_no}행"
            try:
                register(page, item)
                result, fill = "완료", OK_FILL
                ok += 1
                print(f"[{i}/{len(products)}] 완료  {label}")
            except Exception as e:  # 한 건이 실패해도 다음 건은 계속 진행
                reason = str(e).splitlines()[0]
                result, fill = f"실패: {reason}", FAIL_FILL
                fail += 1
                print(f"[{i}/{len(products)}] 실패  {label} - {reason}")
            ws.cell(row=row_no, column=result_col, value=result).fill = fill
            ws.cell(row=row_no, column=time_col, value=datetime.now().strftime("%Y-%m-%d %H:%M:%S"))

        if not args.headless and not args.fast:
            page.wait_for_timeout(2000)  # 마지막 화면을 잠깐 보여 줌
        browser.close()

    OUTPUT_DIR.mkdir(exist_ok=True)
    out = OUTPUT_DIR / f"처리결과_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
    wb.save(out)

    print(f"\n완료 {ok}건, 실패 {fail}건")
    print(f"결과 파일: {out}")


if __name__ == "__main__":
    main()
