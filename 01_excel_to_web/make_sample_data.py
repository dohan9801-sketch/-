"""시연용 상품 목록 엑셀(data/상품목록.xlsx)을 만듭니다.

일부 행에는 일부러 잘못된 값을 넣어, 자동 입력 프로그램이
오류를 어떻게 표시하는지 보여 줄 수 있게 했습니다.
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

HERE = Path(__file__).resolve().parent
OUT = HERE / "data" / "상품목록.xlsx"

HEADERS = ["상품명", "카테고리", "판매가", "재고", "배송방식", "즉시판매", "상품설명"]

ROWS = [
    ["오버핏 코튼 반팔티", "의류", 19900, 120, "택배", "Y", "부드러운 순면 소재의 기본 반팔티"],
    ["와이드 데님 팬츠", "의류", 39800, 45, "택배", "Y", "편안한 와이드 핏 청바지"],
    ["캔버스 에코백", "잡화", 12900, 200, "택배", "Y", "가볍고 튼튼한 데일리 가방"],
    ["가죽 카드지갑", "잡화", 24000, 60, "택배", "N", "소가죽 슬림 카드지갑"],
    ["제주 감귤 3kg", "식품", 18900, 30, "퀵", "Y", "산지 직송 제철 감귤"],
    ["유기농 그래놀라", "식품", 9800, 150, "택배", "Y", "아침 식사용 무가당 그래놀라"],
    ["대용량 수건 5장 세트", "생활용품", 21000, 80, "택배", "Y", "호텔식 40수 코마사 수건"],
    ["스테인리스 텀블러", "생활용품", 15900, 0, "직접수령", "N", "보온·보냉 12시간"],
    ["무선 블루투스 이어폰", "전자기기", 49000, 35, "택배", "Y", "노이즈 캔슬링 지원"],
    ["고속 충전기 25W", "전자기기", 17900, 90, "택배", "Y", "C타입 PD 고속 충전"],
    ["린넨 셔츠", "의류", None, 40, "택배", "Y", "판매가 누락 - 오류 예시"],
    ["접이식 장바구니", "", 6900, 300, "택배", "Y", "카테고리 누락 - 오류 예시"],
]


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "상품목록"
    ws.append(HEADERS)
    for row in ROWS:
        ws.append(row)

    head_fill = PatternFill("solid", fgColor="1F2937")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal="center")
    for row in ws.iter_rows(min_row=2, min_col=3, max_col=3):
        row[0].number_format = "#,##0"
    for col, width in zip("ABCDEFG", [24, 12, 12, 8, 12, 10, 34]):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"

    OUT.parent.mkdir(exist_ok=True)
    wb.save(OUT)
    print(f"만들었습니다: {OUT}")


if __name__ == "__main__":
    main()
